"""Checks on the public analog set, the outputs and the app."""
from io import BytesIO
from pathlib import Path

import opci
import pytest
from openpyxl import load_workbook

from fel1_lng_cost import estimate, load_analogs, make_basis, to_excel_bytes
from fel1_lng_cost import opci_link
from fel1_lng_cost.analogs import check_analogs

APP = Path(__file__).resolve().parent.parent / "app.py"


def test_public_set_is_clean_and_sourced():
    df = load_analogs()
    check = check_analogs(df)
    assert check.ok, check.errors
    assert not check.warnings, check.warnings
    assert df["source_url"].str.startswith("http").all()
    assert (df["notes"].str.len() > 0).all()
    assert (df["unit_cost_reported"] > 0).all()


def test_split_never_exceeds_total():
    df = load_analogs()
    split = df.dropna(subset=["plant_usd_million"])
    assert (split["plant_usd_million"] <= split["cost_usd_million"]).all()


@pytest.mark.parametrize("ftype,size,site", [
    ("onshore_liquefaction", 5.0, "greenfield"),
    ("onshore_liquefaction", 4.0, "brownfield"),
    ("flng", 2.5, "conversion"),
    ("flng", 2.0, "newbuild"),
    ("onshore_regas", 5.0, "greenfield"),
    ("fsru", 5.0, "newbuild"),
])
def test_every_type_runs(ftype, size, site):
    res = estimate(load_analogs(), make_basis(ftype, train_size_mtpa=size, site_type=site))
    assert res.low <= res.central <= res.high
    assert to_excel_bytes(res)[:2] == b"PK"


def test_liquefaction_default_uses_enough_analogs():
    res = estimate(load_analogs(), make_basis(train_size_mtpa=5.0, trains=2))
    assert res.stats["count"] >= 10


def test_app_runs():
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(APP), default_timeout=120).run()
    assert not at.exception
    assert len(at.metric) == 0  # the estimate date is required before anything runs
    assert any("Enter the estimate date" in w.value for w in at.warning)
    date = [t for t in at.sidebar.text_input if t.label.startswith("Estimate date")][0]
    date.set_value(opci_link.show_period(opci_link.latest_quarter())).run()
    assert not at.exception
    assert len(at.metric) >= 4
    assert any(opci_link.stamp_label() in c.value for c in at.caption)


def test_opci_stamp_on_export():
    label = opci_link.stamp_label()
    assert label == f"OPCI {opci.__version__}, {opci_link.latest_release()}"
    assert opci_link.latest_release().startswith("data-")
    res = estimate(load_analogs(), make_basis(train_size_mtpa=5.0))
    ws = load_workbook(BytesIO(to_excel_bytes(res)))["Summary"]
    values = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=2).value
              for i in range(1, ws.max_row + 1)}
    assert values["Cost index"] == label
    esc = load_workbook(BytesIO(to_excel_bytes(res)))["Escalation"]
    assert esc["A1"].value == "Escalation panel: factor per analog"
    assert len(esc._charts) == 1


def test_no_public_file_names_the_indonesian_site():
    """Worked examples call it "an existing LNG hub" (decided 24 October 2026)."""
    root = Path(__file__).resolve().parent.parent
    files = [p for p in root.rglob("*") if p.is_file() and p.suffix in (".py", ".md", ".csv")
             and ".git" not in p.parts and not p.name.startswith("private_")]
    assert files
    for p in files:
        assert "bon" + "tang" not in p.read_text(encoding="utf-8", errors="ignore").lower(), p


def test_package_setup_matches_requirements():
    tomllib = pytest.importorskip("tomllib")  # Python 3.11 and newer

    root = Path(__file__).resolve().parent.parent
    project = tomllib.loads((root / "pyproject.toml").read_text(encoding="utf-8"))["project"]

    def lines(name):
        return [x.strip() for x in (root / name).read_text(encoding="utf-8").splitlines()
                if x.strip() and not x.startswith(("#", "-r"))]

    runtime = lines("requirements.txt") + lines("requirements/opci.txt")
    assert sorted(project["dependencies"]) == sorted(runtime)
    assert not any(x.startswith("pytest") for x in runtime)  # the test runner is a development need
    assert project["optional-dependencies"]["dev"] == [x for x in lines("requirements-dev.txt")]
